#!/usr/bin/env python3
"""
Corrected v3 reproduction script for the Türkiye 2023 A64
capacity-constrained input-output allocation study.

Raw TÜİK workbooks are intentionally not redistributed in this repository.
See data/README.md for acquisition and provenance instructions.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog


EXPECTED_ORIGINAL_XLS_SHA256 = (
    "572d4e0f3cb4288f62d0249e3d8514a61ec40e5888ea9ac323ca740c9ce2378e"
)
SEED = 7
N_TRIALS = 300
PERTURBATIONS = (0.01, 0.05, 0.10, 0.30, 0.60)
SHOCK_LEVELS = (0.05, 0.10, 0.20)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_text(value) -> str:
    return " ".join(str(value).replace("\n", " ").split())


def load_tuik_xlsx(path: Path):
    """
    Read the exact cells used by corrected v3 from sheet T9.

    Excel coordinates:
      codes: B10:B73
      names: C10:C73
      Z0: D10:BO73
      x0: D86:BO86
      hh: BQ10:BQ73
      fc: BT10:BT73
      gcf: BW10:BW73
      exp: BX10:BX73
    """
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb["T9"]

    codes = [clean_text(ws.cell(r, 2).value) for r in range(10, 74)]
    names = [clean_text(ws.cell(r, 3).value) for r in range(10, 74)]

    Z0 = np.array(
        [[float(ws.cell(r, c).value or 0.0) for c in range(4, 68)] for r in range(10, 74)],
        dtype=float,
    )
    x0 = np.array([float(ws.cell(86, c).value or 0.0) for c in range(4, 68)], dtype=float)
    hh = np.array([float(ws.cell(r, 69).value or 0.0) for r in range(10, 74)], dtype=float)
    fc = np.array([float(ws.cell(r, 72).value or 0.0) for r in range(10, 74)], dtype=float)
    gcf = np.array([float(ws.cell(r, 75).value or 0.0) for r in range(10, 74)], dtype=float)
    exp = np.array([float(ws.cell(r, 76).value or 0.0) for r in range(10, 74)], dtype=float)

    return codes, names, Z0, x0, hh, fc, gcf, exp


class Model:
    def __init__(self, codes, names, Z0, x0, hh, fc, gcf, exp):
        self.codes = codes
        self.names = names
        self.Z0 = Z0
        self.x0 = x0
        self.hh = hh
        self.fc = fc
        self.gcf = gcf
        self.exp = exp
        self.N = len(codes)
        self.J = codes.index("D35")
        self.iF = codes.index("F")

        self.f0 = fc + gcf + exp
        if np.any(x0 <= 0):
            raise RuntimeError("Non-positive baseline output x0.")
        if np.any(self.f0 <= 0):
            bad = [(codes[i], self.f0[i]) for i in np.where(self.f0 <= 0)[0]]
            raise RuntimeError(f"Non-positive baseline final use: {bad}")

        if not np.isclose(
            self.f0[self.J],
            hh[self.J] + exp[self.J],
            rtol=1e-10,
            atol=1e-8,
        ):
            raise RuntimeError("D35 final-use identity failed.")

        self.A = Z0 / x0[None, :]
        self.K = (np.eye(self.N) - self.A) * x0[None, :] / self.f0[:, None]

        if not np.allclose(self.K @ np.ones(self.N), np.ones(self.N), atol=1e-9):
            raise RuntimeError("K @ 1 != 1")

        self.lbH = np.zeros(self.N)
        self.lbH[self.J] = hh[self.J] / self.f0[self.J]

        self.w_mon = self.f0 / self.f0.sum() * self.N

        eJ = np.zeros(self.N)
        eJ[self.J] = 1.0
        self.d_feas = float(np.linalg.solve(self.K, eJ)[self.J])

    def solve(self, d, w=None, z_lb=None):
        w = np.ones(self.N) if w is None else np.asarray(w, dtype=float)
        lb = np.zeros(self.N) if z_lb is None else np.asarray(z_lb, dtype=float)
        ub = np.ones(self.N)
        ub[self.J] = 1.0 - d

        res = linprog(
            -(w @ self.K),
            A_ub=np.vstack([self.K, -self.K]),
            b_ub=np.concatenate([np.ones(self.N), -lb]),
            bounds=[(0.0, ub[i]) for i in range(self.N)],
            method="highs",
        )
        if not res.success:
            raise RuntimeError(res.message)

        y = res.x
        z = self.K @ y
        q = 1.0 - y
        return res, y, z, q

    def solve_maxmin(self, d, z_lb=None, eps=1e-9):
        lb = np.zeros(self.N) if z_lb is None else np.asarray(z_lb, dtype=float)
        ub = np.ones(self.N)
        ub[self.J] = 1.0 - d
        idx = [i for i in range(self.N) if i != self.J]

        bnds = [(0.0, ub[i]) for i in range(self.N)] + [(0.0, 1.0)]
        A_ub = np.vstack(
            [
                np.hstack([self.K, np.zeros((self.N, 1))]),
                np.hstack([-self.K, np.zeros((self.N, 1))]),
                np.hstack([-self.K[idx], np.ones((len(idx), 1))]),
            ]
        )
        b_ub = np.concatenate([np.ones(self.N), -lb, np.zeros(len(idx))])

        c1 = np.zeros(self.N + 1)
        c1[-1] = -1.0
        r1 = linprog(c1, A_ub=A_ub, b_ub=b_ub, bounds=bnds, method="highs")
        if not r1.success:
            raise RuntimeError(r1.message)
        t_star = float(r1.x[-1])

        c2 = np.concatenate([-(np.ones(self.N) @ self.K), [0.0]])
        A2 = np.vstack([A_ub, np.concatenate([np.zeros(self.N), [-1.0]])])
        b2 = np.concatenate([b_ub, [-(t_star - eps)]])

        r2 = linprog(c2, A_ub=A2, b_ub=b2, bounds=bnds, method="highs")
        if not r2.success:
            raise RuntimeError(f"Lexicographic stage 2 failed: {r2.message}")

        y = r2.x[: self.N]
        z = self.K @ y
        q = 1.0 - y
        return r1, r2, y, z, q, t_star

    def solve_maxoutput(self, d, z_lb=None):
        lb = np.zeros(self.N) if z_lb is None else np.asarray(z_lb, dtype=float)
        ub = np.ones(self.N)
        ub[self.J] = 1.0 - d

        res = linprog(
            -self.x0 / self.x0.sum(),
            A_ub=np.vstack([self.K, -self.K]),
            b_ub=np.concatenate([np.ones(self.N), -lb]),
            bounds=[(0.0, ub[i]) for i in range(self.N)],
            method="highs",
        )
        if not res.success:
            raise RuntimeError(res.message)

        y = res.x
        return res, y, self.K @ y, 1.0 - y

    def owl(self, q):
        return 100.0 * (q * self.x0).sum() / self.x0.sum()


def write_csv(path: Path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def run(model: Model, outdir: Path):
    outdir.mkdir(parents=True, exist_ok=True)

    # Regression gate
    _, _, _, q0 = model.solve(0.0)
    if np.max(np.abs(q0)) >= 1e-9:
        raise RuntimeError("d=0 regression gate failed.")

    variants = []
    for d in SHOCK_LEVELS:
        _, y, z, q = model.solve(d)
        nz = sum(
            1
            for i in range(model.N)
            if i != model.J and z[i] < 1 - 1e-6
        )
        variants.append(
            {
                "variant": "residual_final_use",
                "d": d,
                "z_D35": float(z[model.J]),
                "mean_q": float(q.mean()),
                "output_weighted_loss_pct": float(model.owl(q)),
                "monetary_coverage": float((z * model.f0).sum() / model.f0.sum()),
                "non_D35_sectors_below_full_coverage": int(nz),
                "z_F_construction": None,
            }
        )

    for d in SHOCK_LEVELS:
        _, y, z_raw, q = model.solve(d, z_lb=model.lbH)
        z = np.clip(z_raw, 0.0, 1.0)
        nz = sum(
            1
            for i in range(model.N)
            if i != model.J and z[i] < 1 - 1e-6
        )
        variants.append(
            {
                "variant": "household_floor_export_first",
                "d": d,
                "z_D35": float(z[model.J]),
                "mean_q": float(q.mean()),
                "output_weighted_loss_pct": float(model.owl(q)),
                "monetary_coverage": float((z * model.f0).sum() / model.f0.sum()),
                "non_D35_sectors_below_full_coverage": int(nz),
                "z_F_construction": float(z[model.iF]),
            }
        )

    weighting_convention = []
    for label, lb in (("residual", None), ("household_floor", model.lbH)):
        for d in SHOCK_LEVELS:
            _, _, z1, _ = model.solve(d, z_lb=lb)
            _, _, z2, _ = model.solve(d, w=model.w_mon, z_lb=lb)
            weighting_convention.append(
                {
                    "variant": label,
                    "d": d,
                    "max_abs_z_diff": float(np.max(np.abs(z1 - z2))),
                }
            )

    rng = np.random.default_rng(SEED)
    perturbation = []
    for p in PERTURBATIONS:
        counts = {}
        for _ in range(N_TRIALS):
            w = 1.0 + rng.uniform(-p, p, model.N)
            _, _, z, _ = model.solve(0.10, w=w, z_lb=model.lbH)
            marginal = [
                model.codes[i]
                for i in range(model.N)
                if i != model.J and 1e-6 < z[i] < 1 - 1e-6
            ]
            key = marginal[0] if len(marginal) == 1 else ("MULTI" if marginal else "NONE")
            counts[key] = counts.get(key, 0) + 1

        top = sorted(counts.items(), key=lambda t: -t[1])
        perturbation.append(
            {
                "perturbation": p,
                "n_trials": N_TRIALS,
                "shares": {k: v / N_TRIALS for k, v in top},
            }
        )

    comparator = []
    for d in SHOCK_LEVELS:
        _, y_add, z_add_raw, q_add = model.solve(d, z_lb=model.lbH)
        z_add = np.clip(z_add_raw, 0.0, 1.0)

        _, _, y_mm, z_mm_raw, q_mm, t_star = model.solve_maxmin(
            d, z_lb=model.lbH
        )
        z_mm = np.clip(z_mm_raw, 0.0, 1.0)

        _, y_mo, z_mo, q_mo = model.solve_maxoutput(d, z_lb=model.lbH)
        maxout_loss = float(model.owl(q_mo))

        for label, z, q in (
            ("additive_linear", z_add, q_add),
            ("lexicographic_maxmin", z_mm, q_mm),
        ):
            comparator.append(
                {
                    "d": d,
                    "objective": label,
                    "sectors_below_full_coverage_all": int(
                        np.sum(z < 1 - 1e-6)
                    ),
                    "sectors_below_full_coverage_non_D35": int(
                        sum(
                            z[i] < 1 - 1e-6
                            for i in range(model.N)
                            if i != model.J
                        )
                    ),
                    "worst_off_z": float(
                        min(z[i] for i in range(model.N) if i != model.J)
                    ),
                    "output_weighted_loss_pct": float(model.owl(q)),
                    "max_output_reference_loss_pct": maxout_loss,
                }
            )

    sampling = (
        "independent multiplicative perturbation w_i = 1 + U(-p, p), "
        f"numpy default_rng(seed={SEED}), {N_TRIALS} trials per level"
    )

    frozen = {
        "model": (
            "capacity-constrained Leontief allocation with "
            "IIM-style inoperability metrics"
        ),
        "formulation": (
            "dimensionless: K y = z, 0<=y,z<=1, "
            "y_J<=1-d, q=1-y"
        ),
        "baseline_objective": (
            "equal-sector normalized final-demand coverage "
            "(normative, not neutral)"
        ),
        "sensitivity_objective": (
            "monetary-weighted coverage; same additive linear form, "
            "different weighting convention"
        ),
        "framing": (
            "best feasible allocations conditional on static domestic "
            "Leontief relations; not forecasts of realized rationing"
        ),
        "supersedes": (
            "v1 numerically invalid; v2 used bisection threshold and "
            "np.maximum(f0,1e-9)"
        ),
        "d_feas_analytic": model.d_feas,
        "d_feas_definition": (
            "maximum capacity loss at which full coverage of all non-D35 "
            "final demand remains feasible; NOT a universal cut-in threshold"
        ),
        "D35_final_use": {
            "household_pct": float(100 * model.hh[model.J] / model.f0[model.J]),
            "exports_pct": float(100 * model.exp[model.J] / model.f0[model.J]),
        },
        "numerics": {
            "K1_max_err": float(
                np.abs(model.K @ np.ones(model.N) - 1).max()
            ),
            "cond2_K": float(np.linalg.cond(model.K)),
            "rank_K": int(np.linalg.matrix_rank(model.K)),
        },
        "sampling_design": sampling,
        "input_sha256": EXPECTED_ORIGINAL_XLS_SHA256,
        "max_output_reference": (
            "max sum x0_i y_i on the same feasible set; "
            "shows the additive coverage objective does NOT minimise output loss"
        ),
        "variants": variants,
        "weighting_convention": weighting_convention,
        "local_perturbation": perturbation,
        "comparator": comparator,
    }

    with (outdir / "25_frozen_results.json").open("w", encoding="utf-8") as fh:
        json.dump(frozen, fh, indent=2, ensure_ascii=False)

    write_csv(outdir / "24_variants.csv", variants)
    write_csv(outdir / "26_comparator.csv", comparator)

    # Full deterministic sector-level output.
    sector_rows = []

    def add_sector_rows(label, d, y, z, q, t_star=None):
        zc = np.clip(z, 0.0, 1.0)
        for i in range(model.N):
            sector_rows.append(
                {
                    "scenario": label,
                    "d": d,
                    "sector_code": model.codes[i],
                    "sector_name": model.names[i],
                    "y": float(y[i]),
                    "z": float(zc[i]),
                    "q": float(q[i]),
                    "t_star_if_maxmin": "" if t_star is None else float(t_star),
                }
            )

    for d in SHOCK_LEVELS:
        _, y, z, q = model.solve(d)
        add_sector_rows("residual_equal", d, y, z, q)

        _, y, z, q = model.solve(d, w=model.w_mon)
        add_sector_rows("residual_monetary", d, y, z, q)

        _, y, z, q = model.solve(d, z_lb=model.lbH)
        add_sector_rows("household_equal_additive", d, y, z, q)

        _, y, z, q = model.solve(d, w=model.w_mon, z_lb=model.lbH)
        add_sector_rows("household_monetary_additive", d, y, z, q)

        _, _, y, z, q, t_star = model.solve_maxmin(d, z_lb=model.lbH)
        add_sector_rows(
            "household_lexicographic_maxmin", d, y, z, q, t_star=t_star
        )

        _, y, z, q = model.solve_maxoutput(d, z_lb=model.lbH)
        add_sector_rows("household_max_output_reference", d, y, z, q)

    write_csv(outdir / "sector_level_solutions.csv", sector_rows)

    # Analytical leverage table used in the completion audit.
    M = np.linalg.inv(model.K)
    r = M[model.J, :]
    leverage_rows = []
    order = np.argsort(-r)
    for rank, i in enumerate(order, start=1):
        leverage_rows.append(
            {
                "rank_equal": rank,
                "sector_code": model.codes[i],
                "sector_name": model.names[i],
                "r_MJ": float(r[i]),
                "monetary_weight": float(model.w_mon[i]),
                "r_over_monetary_weight": float(r[i] / model.w_mon[i]),
            }
        )
    write_csv(outdir / "D35_capacity_relief_leverage.csv", leverage_rows)

    print(f"d* = {model.d_feas:.12f}")
    print(f"written: {outdir}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input-xlsx",
        type=Path,
        required=True,
        help="TÜİK T9 workbook saved as .xlsx without cell edits.",
    )
    parser.add_argument(
        "--original-xls",
        type=Path,
        default=None,
        help=(
            "Optional original .xls. If provided, its SHA-256 is checked "
            "against the frozen provenance hash."
        ),
    )
    parser.add_argument(
        "--outdir",
        type=Path,
        default=Path("results/generated"),
    )
    args = parser.parse_args()

    if args.original_xls is not None:
        actual = sha256_file(args.original_xls)
        if actual != EXPECTED_ORIGINAL_XLS_SHA256:
            raise RuntimeError(
                f"Unexpected original .xls SHA-256: {actual}"
            )
        print("source hash: PASS")

    data = load_tuik_xlsx(args.input_xlsx)
    model = Model(*data)
    run(model, args.outdir)


if __name__ == "__main__":
    main()
